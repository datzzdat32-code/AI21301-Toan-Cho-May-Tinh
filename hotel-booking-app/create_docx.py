import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_heading_styled(doc, text, level):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Times New Roman'
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0, 53, 128) # Navy Blue
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
    elif level == 2:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(2, 132, 199) # Ocean Blue
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
    elif level == 3:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(30, 41, 59)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    return p

def format_table_headers(table, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Format Header Row
    hdr_cells = table.rows[0].cells
    for i, cell in enumerate(hdr_cells):
        set_cell_background(cell, "003580") # Dark Navy
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(255, 255, 255)
    
    # Format Body Rows
    for row_idx, row in enumerate(table.rows[1:], start=1):
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for cell in row.cells:
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(10)

def main():
    doc = docx.Document()
    
    # Page Margins (Normal 1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles default
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(30, 41, 59)

    # ---------------------------------------------------------
    # TRANG BÌA (COVER PAGE)
    # ---------------------------------------------------------
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_header.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\nTRƯỜNG ĐẠI HỌC / HỌC VIỆN CÔNG NGHỆ THÔNG TIN\nKHOA CÔNG NGHỆ THÔNG TIN\n")
    r1.bold = True
    r1.font.size = Pt(12)
    r1.font.color.rgb = RGBColor(71, 85, 105)

    p_star = doc.add_paragraph()
    p_star.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_star = p_star.add_run("------------------***------------------\n\n\n")
    r_star.font.size = Pt(12)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t1 = p_title.add_run("BÁO CÁO QUẢN LÝ ĐỒ ÁN TỐT NGHIỆP\n(PROJECT MANAGEMENT MATRIX DOCUMENT)\n\n")
    r_t1.bold = True
    r_t1.font.size = Pt(18)
    r_t1.font.color.rgb = RGBColor(0, 53, 128)

    r_t2 = p_title.add_run("ĐỀ TÀI: XÂY DỰNG HỆ THỐNG WEBSITE ĐẶT PHÒNG KHÁCH SẠN VÀ RESORT TRỰC TUYẾN TOÀN QUỐC (VNSTAY.VN)\n\n\n")
    r_t2.bold = True
    r_t2.font.size = Pt(15)
    r_t2.font.color.rgb = RGBColor(2, 132, 199)

    # Info Box Table
    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Chuyên ngành:", "Công nghệ Thông tin (Phát triển Web Fullstack)"),
        ("Sinh viên thực hiện:", "Nguyễn Văn A - MSSV: PH12345"),
        ("Giảng viên hướng dẫn:", "ThS. Trần Văn B"),
        ("Năm hoàn thành:", "2026")
    ]
    for idx, (k, v) in enumerate(info_data):
        cell_k, cell_v = info_table.rows[idx].cells
        cell_k.paragraphs[0].add_run(k).bold = True
        cell_v.paragraphs[0].add_run(v)
        set_cell_background(cell_k, "F1F5F9")
        set_cell_background(cell_v, "FFFFFF")

    doc.add_page_break()

    # ---------------------------------------------------------
    # MỤC LỤC & TỔNG QUAN
    # ---------------------------------------------------------
    add_heading_styled(doc, "CHƯƠNG 1: TỔNG QUAN ĐỀ TÀI & CẤU HÌNH CÔNG NGHỆ CHUYÊN SÂU", 1)
    
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.add_run("1.1 Tính Cấp Thiết Của Đề Tài:\n").bold = True
    p.add_run("Nhu cầu du lịch và đặt phòng trực tuyến tại Việt Nam gia tăng nhanh chóng. Khách hàng đòi hỏi nền tảng đặt phòng phải minh bạch chi phí (đã bao gồm Thuế & Phí), thao tác tìm kiếm mượt mà theo 63 tỉnh thành, nhận mã E-Voucher QR Code tức thì và tuyệt đối an toàn thông tin cá nhân.\n\n")
    
    p.add_run("1.2 Mục Tiêu Sản Phẩm (VnStay.vn):\n").bold = True
    p.add_run("• Phía Khách Hàng: Cung cấp giao diện đặt phòng chuẩn OTA du lịch (Traveloka/Agoda style), bộ lọc đa tiêu chí, tính tiền tự động, áp mã voucher, cổng thanh toán mô phỏng QR (VNPay/MoMo/Visa) và quản lý E-Voucher cá nhân.\n")
    p.add_run("• Phía Quản Trị Viên (Admin): Dashboard thống kê doanh thu real-time, đặt phòng trực tiếp tại quầy (Walk-in), theo dõi trạng thái 🟢 Phòng Trống / 🔴 Đã Thuê, quản lý Mã KH (`usr_...`) và xuất báo cáo CSV/Excel.\n\n")

    p.add_run("1.3 Chi Tiết Cấu Hình Công Nghệ (Tech Stack Specs):\n").bold = True

    # Tech Table
    tech_table = doc.add_table(rows=10, cols=3)
    tech_table.rows[0].cells[0].paragraphs[0].text = "Thành Phần"
    tech_table.rows[0].cells[1].paragraphs[0].text = "Công Nghệ / Thư Viện"
    tech_table.rows[0].cells[2].paragraphs[0].text = "Vai Trò & Mô Tả Chi Tiết"

    tech_data = [
        ("Frontend UI", "React 19.2.8 + Vite 8.3.0", "Xây dựng giao diện ứng dụng Single Page Application (SPA) tốc độ cao"),
        ("Routing", "React Router DOM v7.18.4", "Điều hướng trang mượt mà không reload (HomePage, Rooms, Checkout, Admin)"),
        ("CSS Framework", "Tailwind CSS v4.3.3", "Thiết kế giao diện responsive Trắng - Xanh chuẩn du lịch"),
        ("Icons & HTTP", "Lucide React + Axios v1.20.0", "Bộ icon thương mại trực quan & xử lý HTTP REST API"),
        ("Backend Server", "Node.js + Express.js v5.2.1", "Xây dựng hệ thống máy chủ RESTful API Cổng 5000"),
        ("Database Engine", "SQLite3 v6.0.1", "Cơ sở dữ liệu quan hệ nhúng độc lập tại backend/hotel.db (6 bảng)"),
        ("Authentication", "JsonWebToken (JWT) v9.0.3", "Mã định danh xác thực phiên làm việc an toàn 24h"),
        ("Security Hashing", "Bcrypt.js v3.0.3", "Mã hóa băm mật khẩu chuẩn Salt round 10"),
        ("Security Protection", "Helmet v8.3.0 + Rate-Limit v8.7.0", "Bảo vệ HTTP Security Headers & khóa IP 15 phút nếu nhập sai quá 5 lần")
    ]

    for idx, row in enumerate(tech_data, start=1):
        cells = tech_table.rows[idx].cells
        cells[0].paragraphs[0].text = row[0]
        cells[1].paragraphs[0].text = row[1]
        cells[2].paragraphs[0].text = row[2]

    format_table_headers(tech_table)

    # ---------------------------------------------------------
    # CHƯƠNG 2: MA TRẬN WBS
    # ---------------------------------------------------------
    add_heading_styled(doc, "CHƯƠNG 2: MA TRẬN PHÂN CÔNG NHIỆM VỤ & TIẾN ĐỘ THI CÔNG (WBS)", 1)
    
    wbs_table = doc.add_table(rows=19, cols=6)
    headers = ["Mã WBS", "Hạng Mục Công Việc Chi Tiết", "Người Thực Hiện", "Thời Gian", "Sản Phẩm Đầu Ra", "Trạng Thái"]
    for i, h in enumerate(headers):
        wbs_table.rows[0].cells[i].paragraphs[0].text = h

    wbs_data = [
        ("WBS 1.1", "Nghiên cứu UX/UI du lịch từ Traveloka/Agoda", "PM / UI Lead", "Tuần 1", "Tài liệu Yêu cầu", "🟢 PASS 100%"),
        ("WBS 1.2", "Thiết kế CSDL SQLite 6 bảng thực thể", "DB Dev", "Tuần 1", "backend/database.js", "🟢 PASS 100%"),
        ("WBS 1.3", "Vẽ sơ đồ ERD & mã DBML dbdiagram.io", "DB Dev", "Tuần 1", "Sơ đồ ERD & DBML", "🟢 PASS 100%"),
        ("WBS 2.1", "Khởi tạo Express Server Cổng 5000 kết nối SQLite", "Backend Dev", "Tuần 2", "backend/server.js", "🟢 PASS 100%"),
        ("WBS 2.2", "Viết API Đăng ký / Đăng nhập (Bcrypt + JWT)", "Backend Dev", "Tuần 2", "Auth API Router", "🟢 PASS 100%"),
        ("WBS 2.3", "Viết API Khách sạn & Tìm kiếm 63 tỉnh thành", "Backend Dev", "Tuần 2", "Rooms API Router", "🟢 PASS 100%"),
        ("WBS 2.4", "Viết API Đặt phòng, Thanh toán QR & Hủy phòng", "Backend Dev", "Tuần 3", "Bookings Router", "🟢 PASS 100%"),
        ("WBS 2.5", "Viết API Admin Stats, Walk-in & Maintenance", "Backend Dev", "Tuần 3", "Admin Router", "🟢 PASS 100%"),
        ("WBS 2.6", "Bảo mật Helmet, Rate-limit 5/15p & RBAC Admin", "Security Dev", "Tuần 3", "Security Middlewares", "🟢 PASS 100%"),
        ("WBS 3.1", "Setup React 19 + Tailwind CSS Trắng - Xanh", "Frontend Dev", "Tuần 4", "Base Frontend App", "🟢 PASS 100%"),
        ("WBS 3.2", "Dựng Trang Chủ (Hero search, Vouchers, Reviews)", "Frontend Dev", "Tuần 4", "Page HomePage.jsx", "🟢 PASS 100%"),
        ("WBS 3.3", "Dựng Trang Danh Sách & Bộ lọc đa tiêu chí", "Frontend Dev", "Tuần 4", "Page RoomsPage.jsx", "🟢 PASS 100%"),
        ("WBS 3.4", "Dựng Trang Chi Tiết Phòng & Form Review", "Frontend Dev", "Tuần 5", "Page RoomDetailPage.jsx", "🟢 PASS 100%"),
        ("WBS 3.5", "Dựng Trang Checkout & Modal Quét QR Code", "Frontend Dev", "Tuần 5", "CheckoutPage & PaymentModal", "🟢 PASS 100%"),
        ("WBS 3.6", "Dựng Trang Đơn Của Tôi & Tra cứu E-Voucher QR", "Frontend Dev", "Tuần 5", "Page MyBookingsPage.jsx", "🟢 PASS 100%"),
        ("WBS 3.7", "Dựng Dashboard Admin (Stats, Walk-in, CSV)", "Frontend Dev", "Tuần 6", "Page AdminDashboardPage.jsx", "🟢 PASS 100%"),
        ("WBS 4.1", "Tối ưu hóa UI/UX loại bỏ cảm giác AI", "UI Lead", "Tuần 7", "Polish UI Components", "🟢 PASS 100%"),
        ("WBS 4.2", "Đẩy mã nguồn lên GitHub & Viết Báo cáo DATN", "DevOps / PM", "Tuần 7", "GitHub Repository & Doc", "🟢 PASS 100%")
    ]

    for idx, row in enumerate(wbs_data, start=1):
        cells = wbs_table.rows[idx].cells
        for i, val in enumerate(row):
            cells[i].paragraphs[0].text = val

    format_table_headers(wbs_table)

    # ---------------------------------------------------------
    # CHƯƠNG 3: YÊU CẦU CHỨC NĂNG
    # ---------------------------------------------------------
    add_heading_styled(doc, "CHƯƠNG 3: MA TRẬN YÊU CẦU CHỨC NĂNG & PHI CHỨC NĂNG", 1)

    add_heading_styled(doc, "3.1 Nhóm Chức Năng Phía Khách Hàng (Customer Features)", 2)
    
    req_cust_table = doc.add_table(rows=8, cols=3)
    req_cust_table.rows[0].cells[0].paragraphs[0].text = "Mã Yêu Cầu"
    req_cust_table.rows[0].cells[1].paragraphs[0].text = "Tên Chức Năng"
    req_cust_table.rows[0].cells[2].paragraphs[0].text = "Mô Tả Kỹ Thuật & Tiêu Chí Nghiệm Thu"

    cust_data = [
        ("F-C01", "Tìm kiếm địa điểm 63 tỉnh thành", "Tìm phòng theo Đà Nẵng, Phú Quốc, Đà Lạt, Nha Trang,... ngày nhận/trả & số khách."),
        ("F-C02", "Bộ lọc đa tiêu chí", "Lọc theo khoảng giá, hạng phòng (Standard, Deluxe, Suite, Villa), sắp xếp giá tăng/giảm."),
        ("F-C03", "Thẻ phòng chuẩn OTA", "Hiển thị điểm 10 (9.4/10), giá cũ bị gạch, tag 'Miễn phí hủy phòng', 'Bao gồm ăn sáng'."),
        ("F-C04", "Chi tiết phòng & Gallery", "Xem ảnh carousel, m², số giường, chính sách nhận (14:00) / trả (12:00) phòng."),
        ("F-C05", "Thanh toán QR mô phỏng", "Quét mã QR VNPay / MoMo / Thẻ Visa / Tiền mặt kèm đồng hồ đếm ngược 15 phút."),
        ("F-C06", "Vé điện tử E-Voucher QR", "Tạo mã QR E-Voucher duy nhất cho từng đơn hàng thành công để nhận phòng."),
        ("F-C07", "Hủy đơn trực tuyến", "Cho phép khách hàng chủ động hủy đơn trong trang cá nhân theo quy định.")
    ]

    for idx, row in enumerate(cust_data, start=1):
        cells = req_cust_table.rows[idx].cells
        cells[0].paragraphs[0].text = row[0]
        cells[1].paragraphs[0].text = row[1]
        cells[2].paragraphs[0].text = row[2]

    format_table_headers(req_cust_table)

    add_heading_styled(doc, "3.2 Nhóm Chức Năng Phía Quản Trị Viên (Admin Features)", 2)

    req_admin_table = doc.add_table(rows=8, cols=3)
    req_admin_table.rows[0].cells[0].paragraphs[0].text = "Mã Yêu Cầu"
    req_admin_table.rows[0].cells[1].paragraphs[0].text = "Tên Chức Năng"
    req_admin_table.rows[0].cells[2].paragraphs[0].text = "Mô Tả Kỹ Thuật & Tiêu Chí Nghiệm Thu"

    admin_data = [
        ("F-A01", "Thống kê doanh thu Real-time", "Báo cáo tổng doanh số thực tế, doanh số chờ duyệt, tỷ lệ lấp đầy phòng (%) & kênh thanh toán."),
        ("F-A02", "Đặt phòng tại quầy (Walk-in)", "Tạo đơn trực tiếp cho khách vãng lai nhận phòng ngay tại lễ tân mà không cần tài khoản."),
        ("F-A03", "Theo dõi trạng thái phòng", "Phân loại 🟢 Phòng Trống vs 🔴 Đã Thuê (Hiển thị Tên khách, SĐT & thời gian ở)."),
        ("F-A04", "Bật/Tắt bảo trì phòng", "Nút chuyển trạng thái phòng giữa 'Sẵn sàng đón khách' và 'Đang bảo trì'."),
        ("F-A05", "Quản lý Mã Khách Hàng", "Tra cứu Mã KH (`usr_...`), tổng chi tiêu tích lũy & danh sách phòng từng khách đã thuê."),
        ("F-A06", "Xuất báo cáo CSV/Excel", "Nút 'Xuất Báo Cáo CSV' tự động tải file `.csv` dữ liệu đơn hàng mở được trên Excel."),
        ("F-A07", "Tìm kiếm toàn cục Admin", "Ô tìm kiếm nhanh theo số phòng (101, 102), Mã KH, SĐT, Tên khách hoặc Mã đơn.")
    ]

    for idx, row in enumerate(admin_data, start=1):
        cells = req_admin_table.rows[idx].cells
        cells[0].paragraphs[0].text = row[0]
        cells[1].paragraphs[0].text = row[1]
        cells[2].paragraphs[0].text = row[2]

    format_table_headers(req_admin_table)

    # ---------------------------------------------------------
    # CHƯƠNG 4: CSDL & API
    # ---------------------------------------------------------
    add_heading_styled(doc, "CHƯƠNG 4: THIẾT KẾ CƠ SỞ DỮ LIỆU & RESTFUL API SPECIFICATIONS", 1)

    add_heading_styled(doc, "4.1 Cấu Trúc 6 Bảng Cơ Sở Dữ Liệu SQLite (hotel.db)", 2)

    db_summary_table = doc.add_table(rows=7, cols=3)
    db_summary_table.rows[0].cells[0].paragraphs[0].text = "Tên Bảng (Table)"
    db_summary_table.rows[0].cells[1].paragraphs[0].text = "Khóa Chính (PK) & Khóa Ngoại (FK)"
    db_summary_table.rows[0].cells[2].paragraphs[0].text = "Chức Năng Lưu Trữ"

    db_rows = [
        ("users", "PK: id (usr_...)", "Lưu tài khoản Khách hàng & Admin, mật khẩu băm Bcrypt, Email (Unique), SĐT"),
        ("room_types", "PK: id (rt_...)", "Phân loại hạng phòng tiêu chuẩn (Standard, Deluxe, Suite, Villa)"),
        ("rooms", "PK: id (rm_...), FK: room_type_id", "Lưu thông tin phòng, số phòng (Unique: 101, 102), địa điểm, giá/đêm, JSON ảnh & tiện nghi"),
        ("bookings", "PK: id (bk_...), FK: user_id, room_id", "Lưu thông tin đơn đặt phòng, ngày Check-in/out, tổng tiền, trạng thái (PENDING, CONFIRMED)"),
        ("payments", "PK: id (pay_...), FK: booking_id", "Lưu lịch sử giao dịch thanh toán tài chính (VNPAY, MOMO, CARD, CASH)"),
        ("reviews", "PK: id (rev_...), FK: user_id, room_id", "Lưu bình luận & điểm sao đánh giá ($1 \\rightarrow 5$ sao) cho từng phòng")
    ]

    for idx, row in enumerate(db_rows, start=1):
        cells = db_summary_table.rows[idx].cells
        cells[0].paragraphs[0].text = row[0]
        cells[1].paragraphs[0].text = row[1]
        cells[2].paragraphs[0].text = row[2]

    format_table_headers(db_summary_table)

    add_heading_styled(doc, "4.2 Danh Sách RESTful API Endpoints", 2)

    api_table = doc.add_table(rows=9, cols=4)
    api_table.rows[0].cells[0].paragraphs[0].text = "Method"
    api_table.rows[0].cells[1].paragraphs[0].text = "Endpoint API"
    api_table.rows[0].cells[2].paragraphs[0].text = "Quyền"
    api_table.rows[0].cells[3].paragraphs[0].text = "Mô Tả Kỹ Thuật"

    api_rows = [
        ("POST", "/api/auth/login", "Public", "Đăng nhập hệ thống (Bảo vệ bởi Rate Limiter 5 lần/15 phút)"),
        ("GET", "/api/rooms", "Public", "Lấy danh sách tất cả phòng & lọc theo 63 tỉnh thành, giá tiền"),
        ("GET", "/api/rooms/:id", "Public", "Lấy thông tin chi tiết 1 phòng & các đánh giá khách hàng"),
        ("POST", "/api/bookings", "Customer", "Tạo đơn đặt phòng mới & chọn cổng thanh toán QR"),
        ("GET", "/api/bookings/my-bookings", "Customer", "Lấy danh sách đơn hàng cá nhân & xem QR E-Voucher"),
        ("GET", "/api/admin/stats", "Admin Only", "Lấy thống kê tổng quan doanh thu, tỷ lệ phòng lấp đầy"),
        ("POST", "/api/admin/bookings/walk-in", "Admin Only", "Tạo đơn đặt phòng trực tiếp tại quầy cho khách vãng lai"),
        ("PUT", "/api/admin/rooms/:id/toggle-maintenance", "Admin Only", "Bật/Tắt trạng thái bảo trì phòng (🟢 Sẵn sàng / 🔴 Bảo trì)")
    ]

    for idx, row in enumerate(api_rows, start=1):
        cells = api_table.rows[idx].cells
        cells[0].paragraphs[0].text = row[0]
        cells[1].paragraphs[0].text = row[1]
        cells[2].paragraphs[0].text = row[2]
        cells[3].paragraphs[0].text = row[3]

    format_table_headers(api_table)

    # ---------------------------------------------------------
    # CHƯƠNG 5: TEST CASES
    # ---------------------------------------------------------
    add_heading_styled(doc, "CHƯƠNG 5: BẢNG MA TRẬN KỊCH BẢN KIỂM THỬ HỆ THỐNG (TEST CASES)", 1)

    tc_table = doc.add_table(rows=7, cols=5)
    tc_table.rows[0].cells[0].paragraphs[0].text = "Mã TC"
    tc_table.rows[0].cells[1].paragraphs[0].text = "Kịch Bản Kiểm Thử"
    tc_table.rows[0].cells[2].paragraphs[0].text = "Các Bước Thực Hiện"
    tc_table.rows[0].cells[3].paragraphs[0].text = "Kết Quả Kỳ Vọng"
    tc_table.rows[0].cells[4].paragraphs[0].text = "Trạng Thái"

    tc_rows = [
        ("TC-01", "Kiểm thử khóa Brute-force Login", "Nhập sai mật khẩu Admin 5 lần liên tiếp", "Trả về lỗi HTTP 429 & khóa IP tạm thời 15 phút", "🟢 PASS"),
        ("TC-02", "Kiểm thử Phân quyền JWT Admin", "Khách hàng mở URL /admin hoặc API admin", "Hệ thống chặn & trả về lỗi 403 Forbidden", "🟢 PASS"),
        ("TC-03", "Kiểm thử Tìm kiếm 63 tỉnh thành", "Chọn địa điểm Đà Nẵng tại ô tìm kiếm", "Trả về đúng danh sách phòng tại Đà Nẵng", "🟢 PASS"),
        ("TC-04", "Kiểm thử Đặt phòng & E-Voucher QR", "Chọn ngày Check-in/out -> Quét QR VNPay", "Tạo đơn thành công & hiển thị E-Voucher QR", "🟢 PASS"),
        ("TC-05", "Kiểm thử Walk-in Booking tại quầy", "Admin bấm '➕ Đặt Phòng Tại Quầy' -> Chọn phòng", "Tạo đơn CONFIRMED mà khách không cần tài khoản", "🟢 PASS"),
        ("TC-06", "Kiểm thử Xuất Báo Cáo CSV", "Admin bấm '📥 Xuất Báo Cáo CSV'", "Tải xuống file VnStay_BaoCaoDoanhThu.csv", "🟢 PASS")
    ]

    for idx, row in enumerate(tc_rows, start=1):
        cells = tc_table.rows[idx].cells
        for i, val in enumerate(row):
            cells[i].paragraphs[0].text = val

    format_table_headers(tc_table)

    # Save document
    doc_path = r"c:\Users\PC\OneDrive\Desktop\AI21301 - Toán cho học máy\hotel-booking-app\DATN_VnStay_BaoCaoQuanLyDuAn.docx"
    doc.save(doc_path)
    print("SUCCESS: File generated successfully!")

if __name__ == '__main__':
    main()
